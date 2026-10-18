"""Extract plain text from an uploaded resume, choosing the method by file type.

| File                         | Method                                   |
|------------------------------|------------------------------------------|
| PDF with a text layer        | PyMuPDF reads the text directly          |
| Scanned PDF (page is a photo)| PyMuPDF renders the page, RapidOCR reads |
| DOCX                         | python-docx reads paragraphs and tables  |
| JPG / PNG / WEBP             | RapidOCR                                 |

The file type is detected from the file's content (magic bytes), never from
its name, so a renamed file cannot pretend to be a resume.
"""

from __future__ import annotations

import io
import logging
import re
import threading
import zipfile
from dataclasses import dataclass, field
from enum import StrEnum

MAX_FILE_BYTES = 10 * 1024 * 1024  # requirements §11
MAX_PAGES = 10                     # resumes are short; protects the server
MIN_TEXT_CHARS_PER_PAGE = 40       # less than this on a PDF page = scanned page
OCR_RENDER_DPI = 200
MAX_IMAGE_SIDE = 3000              # larger photos are scaled down before OCR


class FileKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    JPEG = "jpeg"
    PNG = "png"
    WEBP = "webp"


class UnsupportedFile(ValueError):
    """The file cannot be read as a resume. The message is safe to show users."""


@dataclass(frozen=True)
class ExtractedText:
    text: str
    kind: FileKind
    method: str                     # pdf-text | pdf-ocr | pdf-mixed | docx | image-ocr
    pages: int = 1
    ocr_pages: int = 0
    ocr_confidence: float | None = None  # average OCR line score, 0–1
    warnings: list[str] = field(default_factory=list)


def detect_kind(data: bytes) -> FileKind:
    """Identify the file type from its first bytes."""
    if data.startswith(b"%PDF-"):
        return FileKind.PDF
    if data.startswith(b"\xff\xd8\xff"):
        return FileKind.JPEG
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return FileKind.PNG
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return FileKind.WEBP
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return FileKind.DOCX
        except zipfile.BadZipFile:
            pass
    raise UnsupportedFile("Unsupported file type. Please upload a PDF, DOCX, JPG or PNG file.")


def extract_text(data: bytes) -> ExtractedText:
    """Read a resume file and return its text plus how it was read."""
    if not data:
        raise UnsupportedFile("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise UnsupportedFile("The file is larger than 10 MB.")

    kind = detect_kind(data)
    if kind is FileKind.PDF:
        result = _from_pdf(data)
    elif kind is FileKind.DOCX:
        result = _from_docx(data)
    else:
        result = _from_image(data, kind)

    if not result.text.strip():
        raise UnsupportedFile("No text could be found in this file. Please upload a clearer copy or enter your details manually.")
    return result


# ---------------------------------------------------------------- PDF

def _from_pdf(data: bytes) -> ExtractedText:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise UnsupportedFile("This PDF could not be opened. It may be damaged.") from exc
    with doc:
        if doc.needs_pass:
            raise UnsupportedFile("This PDF is password-protected. Please upload a copy without a password.")

        warnings = []
        page_count = doc.page_count
        if page_count > MAX_PAGES:
            warnings.append(f"Only the first {MAX_PAGES} of {page_count} pages were read.")

        texts, scores, ocr_pages = [], [], 0
        for page in list(doc)[:MAX_PAGES]:
            text = page.get_text("text")
            if len(text.strip()) >= MIN_TEXT_CHARS_PER_PAGE:
                texts.append(text)
                continue
            # Scanned page: render it as an image and read it with OCR.
            png = page.get_pixmap(dpi=OCR_RENDER_DPI).tobytes("png")
            ocr_text, ocr_scores = _ocr(_image_array(png))
            texts.append(ocr_text)
            scores.extend(ocr_scores)
            ocr_pages += 1

    read_pages = min(page_count, MAX_PAGES)
    method = "pdf-text" if ocr_pages == 0 else "pdf-ocr" if ocr_pages == read_pages else "pdf-mixed"
    return ExtractedText(
        text=_clean("\n\n".join(texts)),
        kind=FileKind.PDF,
        method=method,
        pages=read_pages,
        ocr_pages=ocr_pages,
        ocr_confidence=_mean(scores),
        warnings=warnings,
    )


# ---------------------------------------------------------------- DOCX

def _from_docx(data: bytes) -> ExtractedText:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedFile("This Word document could not be opened. It may be damaged.") from exc

    parts = [p.text for p in document.paragraphs]
    # Many resume templates put skills or dates in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(dict.fromkeys(c for c in cells if c)))  # merged cells repeat; keep once
    return ExtractedText(text=_clean("\n".join(parts)), kind=FileKind.DOCX, method="docx")


# ---------------------------------------------------------------- Images / OCR

def _from_image(data: bytes, kind: FileKind) -> ExtractedText:
    try:
        image = _image_array(data)
    except Exception as exc:
        raise UnsupportedFile("This image could not be opened. It may be damaged.") from exc
    text, scores = _ocr(image)
    warnings = []
    confidence = _mean(scores)
    if confidence is not None and confidence < 0.8:
        warnings.append("The image is hard to read. Please check the extracted text, or upload a clearer copy.")
    return ExtractedText(
        text=_clean(text), kind=kind, method="image-ocr", ocr_pages=1, ocr_confidence=confidence, warnings=warnings
    )


def _image_array(data: bytes):
    import numpy as np
    from PIL import Image, ImageOps

    with Image.open(io.BytesIO(data)) as img:
        img = ImageOps.exif_transpose(img)  # phone photos may be stored rotated
        img = img.convert("RGB")
        img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
        return np.array(img)


_ocr_engine = None
_ocr_lock = threading.Lock()


def _get_ocr():
    """Load the OCR models once (about 1 s) and reuse them."""
    global _ocr_engine
    with _ocr_lock:
        if _ocr_engine is None:
            from rapidocr import RapidOCR

            logging.getLogger("RapidOCR").setLevel(logging.WARNING)
            _ocr_engine = RapidOCR()
        return _ocr_engine


def _ocr(image) -> tuple[str, list[float]]:
    result = _get_ocr()(image)
    if not result.txts:
        return "", []
    return "\n".join(result.txts), [float(s) for s in result.scores]


# ---------------------------------------------------------------- helpers

def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _mean(values: list[float]) -> float | None:
    return round(sum(values) / len(values), 3) if values else None
