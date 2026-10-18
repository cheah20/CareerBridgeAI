"""Resume text extraction: each file type uses the right reader."""

import io

import docx
import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont

from app.services.resume_text import (
    MAX_FILE_BYTES,
    FileKind,
    UnsupportedFile,
    detect_kind,
    extract_text,
)

LINES = [
    "AINA RAHMAN",
    "Data Analyst, Kuala Lumpur",
    "Skills: Python, SQL, Power BI",
    "Education: BSc Computer Science, UTeM",
]


def make_png(lines=LINES) -> bytes:
    font = ImageFont.load_default(size=30)
    img = Image.new("RGB", (900, 70 * len(lines) + 40), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((30, 30 + i * 70), line, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def make_text_pdf(lines=LINES) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + i * 20), line, fontsize=12)
    return doc.tobytes()


def make_scanned_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(36, 36, 560, 300), stream=make_png())
    return doc.tobytes()


def make_docx() -> bytes:
    document = docx.Document()
    for line in LINES[:2]:
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, SQL, Power BI"
    table.cell(1, 0).text = "Education"
    table.cell(1, 1).text = "BSc Computer Science, UTeM"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ---- file type detection

def test_detects_type_from_content_not_name():
    assert detect_kind(make_text_pdf()) is FileKind.PDF
    assert detect_kind(make_docx()) is FileKind.DOCX
    assert detect_kind(make_png()) is FileKind.PNG


@pytest.mark.parametrize(
    "data",
    [
        b"MZ\x90\x00 this is a Windows program",          # .exe renamed to .pdf
        b"PK\x03\x04" + b"\x00" * 40,                      # broken zip
        b"plain text pretending to be a resume",
    ],
)
def test_rejects_unsupported_files(data):
    with pytest.raises(UnsupportedFile):
        extract_text(data)


def test_rejects_empty_and_oversized_files():
    with pytest.raises(UnsupportedFile, match="empty"):
        extract_text(b"")
    with pytest.raises(UnsupportedFile, match="10 MB"):
        extract_text(b"%PDF-" + b"0" * MAX_FILE_BYTES)


def test_rejects_password_protected_pdf():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret resume " * 10)
    locked = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="owner")
    with pytest.raises(UnsupportedFile, match="password"):
        extract_text(locked)


# ---- readers

def test_text_pdf_is_read_directly_without_ocr():
    result = extract_text(make_text_pdf())
    assert result.method == "pdf-text"
    assert result.ocr_pages == 0
    assert "Skills: Python, SQL, Power BI" in result.text


def test_docx_reads_paragraphs_and_tables():
    result = extract_text(make_docx())
    assert result.method == "docx"
    assert "AINA RAHMAN" in result.text
    assert "Skills | Python, SQL, Power BI" in result.text


def test_image_is_read_with_ocr():
    result = extract_text(make_png())
    assert result.method == "image-ocr"
    assert result.ocr_confidence is not None and result.ocr_confidence > 0.8
    text = result.text.lower()
    for word in ("python", "sql", "utem"):
        assert word in text


def test_scanned_pdf_falls_back_to_ocr():
    result = extract_text(make_scanned_pdf())
    assert result.method == "pdf-ocr"
    assert result.ocr_pages == 1
    assert "python" in result.text.lower()


def test_pdf_with_text_and_scanned_pages_is_mixed():
    doc = pymupdf.open(stream=make_text_pdf(), filetype="pdf")
    doc.insert_pdf(pymupdf.open(stream=make_scanned_pdf(), filetype="pdf"))
    result = extract_text(doc.tobytes())
    assert result.method == "pdf-mixed"
    assert (result.pages, result.ocr_pages) == (2, 1)


def test_blank_document_explains_there_is_no_text():
    blank = pymupdf.open()
    blank.new_page()
    with pytest.raises(UnsupportedFile, match="No text"):
        extract_text(blank.tobytes())
